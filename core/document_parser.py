from pathlib import Path
from typing import Optional
import pypdf
import docx
from pptx import Presentation


class DocumentParser:
    SUPPORTED_FORMATS = ['.pdf', '.docx', '.pptx', '.txt', '.md']

    @staticmethod
    def parse(file_path: str) -> str:
        ext = Path(file_path).suffix.lower()
        
        if ext == '.pdf':
            return DocumentParser.parse_pdf(file_path)
        elif ext == '.docx':
            return DocumentParser.parse_docx(file_path)
        elif ext == '.pptx':
            return DocumentParser.parse_pptx(file_path)
        elif ext in ['.txt', '.md']:
            return DocumentParser.parse_txt(file_path)
        else:
            raise ValueError(f"不支持的文件格式: {ext}")

    @staticmethod
    def parse_pdf(file_path: str) -> str:
        text = ""
        with open(file_path, 'rb') as f:
            reader = pypdf.PdfReader(f)
            for page in reader.pages:
                text += page.extract_text() + "\n"
        return text

    @staticmethod
    def parse_docx(file_path: str) -> str:
        doc = docx.Document(file_path)
        return "\n".join([para.text for para in doc.paragraphs])

    @staticmethod
    def parse_pptx(file_path: str) -> str:
        prs = Presentation(file_path)
        text = ""
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
        return text

    @staticmethod
    def parse_txt(file_path: str) -> str:
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()